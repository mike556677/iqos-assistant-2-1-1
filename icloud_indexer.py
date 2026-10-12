#!/usr/bin/env python3
"""Read-only iCloud Drive text indexer for IQOS Assistant 2.1.1.

Scans only the folder explicitly supplied with --folder. It never renames,
moves, deletes, or uploads original files. By default it performs a local
preview only; --upload sends extracted text + metadata over HTTPS to the
configured IQOS API.

Expected API contract for the later backend step:
POST {IQOS_API_URL}/icloud/index
Header: X-IQOS-API-Key: <secret>
JSON: {"documents": [{"path", "name", "extension", "size_bytes",
"modified_at", "text"}, ...]}
"""
from __future__ import annotations

import argparse
import html
import json
import os
import sys
import urllib.error
import urllib.request
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARS = 100_000
BATCH_SIZE = 10
TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".html", ".htm", ".xml", ".log"}
OPTIONAL_EXTENSIONS = {".pdf", ".docx"}
# Windows cloud placeholder flags: avoid opening files that may not be present locally.
FILE_ATTRIBUTE_RECALL_ON_OPEN = 0x00040000
FILE_ATTRIBUTE_RECALL_ON_DATA_ACCESS = 0x00400000


def is_cloud_placeholder(path: Path) -> bool:
    try:
        attrs = getattr(path.stat(), "st_file_attributes", 0)
        return bool(attrs & (FILE_ATTRIBUTE_RECALL_ON_OPEN | FILE_ATTRIBUTE_RECALL_ON_DATA_ACCESS))
    except OSError:
        return True


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in TEXT_EXTENSIONS:
        raw = path.read_bytes()
        # Support common text encodings without failing the whole scan.
        for encoding in ("utf-8-sig", "utf-16", "cp1252"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            text = raw.decode("utf-8", errors="replace")
        if suffix in {".html", ".htm"}:
            import re
            text = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", text)
            text = re.sub(r"(?s)<[^>]+>", " ", text)
            text = html.unescape(text)
        return text[:MAX_TEXT_CHARS]
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("PDF support needs: py -m pip install pypdf") from exc
        reader = PdfReader(str(path))
        parts = []
        total = 0
        for page in reader.pages:
            part = page.extract_text() or ""
            parts.append(part)
            total += len(part)
            if total >= MAX_TEXT_CHARS:
                break
        return "\n".join(parts)[:MAX_TEXT_CHARS]
    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("DOCX support needs: py -m pip install python-docx") from exc
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)[:MAX_TEXT_CHARS]
    return ""


def document_record(path: Path, root: Path) -> dict[str, Any] | None:
    suffix = path.suffix.lower()
    if suffix not in TEXT_EXTENSIONS | OPTIONAL_EXTENSIONS:
        return None
    try:
        stat = path.stat()
    except OSError as exc:
        print(f"SKIP (cannot inspect): {path.name} ({exc})")
        return None
    if stat.st_size > MAX_FILE_BYTES:
        print(f"SKIP (over 10 MB): {path.name}")
        return None
    if is_cloud_placeholder(path):
        print(f"SKIP (iCloud file not fully local): {path.name}")
        return None
    try:
        text = extract_text(path)
    except (OSError, RuntimeError, ValueError, json.JSONDecodeError) as exc:
        print(f"SKIP (text extraction failed): {path.name} ({exc})")
        return None
    if not text.strip():
        print(f"SKIP (no extractable text): {path.name}")
        return None
    try:
        relative_path = str(path.relative_to(root))
    except ValueError:
        relative_path = path.name
    return {
        "path": relative_path,
        "name": path.name,
        "extension": suffix,
        "size_bytes": stat.st_size,
        "modified_at": datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc).isoformat(),
        "text": text,
    }


def send_batch(api_url: str, api_key: str, documents: list[dict[str, Any]]) -> None:
    base = api_url.rstrip("/")
    if not base.lower().startswith("https://"):
        raise ValueError("IQOS_API_URL must start with https://. Refusing to send document text over plain HTTP.")
    url = base + "/icloud/index"
    payload = json.dumps({"documents": documents}, ensure_ascii=False).encode("utf-8")
    request = urllib.request.Request(
        url,
        data=payload,
        method="POST",
        headers={
            "Content-Type": "application/json; charset=utf-8",
            "X-IQOS-API-Key": api_key,
            "User-Agent": "IQOS-Assistant-2.1.1-iCloudIndexer/1.0",
        },
    )
    try:
        with urllib.request.urlopen(request, timeout=45) as response:
            body = response.read(4000).decode("utf-8", errors="replace")
            if not 200 <= response.status < 300:
                raise RuntimeError(f"API returned HTTP {response.status}: {body}")
            print(f"Uploaded metadata + extracted text for {len(documents)} file(s). API response: {body}")
    except urllib.error.HTTPError as exc:
        body = exc.read(4000).decode("utf-8", errors="replace")
        raise RuntimeError(f"API returned HTTP {exc.code}: {body}") from exc
    except urllib.error.URLError as exc:
        raise RuntimeError(f"Could not reach the IQOS API: {exc.reason}") from exc


def main() -> int:
    parser = argparse.ArgumentParser(description="Read-only iCloud Drive indexer for IQOS Assistant 2.1.1")
    parser.add_argument("--folder", required=True, help="Exact local path to the dedicated iCloud Drive folder")
    parser.add_argument("--upload", action="store_true", help="Send extracted text + metadata to the secured IQOS API")
    args = parser.parse_args()

    root = Path(os.path.expandvars(os.path.expanduser(args.folder))).resolve()
    if not root.exists() or not root.is_dir():
        print(f"ERROR: folder does not exist or is not a folder: {root}", file=sys.stderr)
        return 2

    print(f"Read-only scan folder: {root}")
    print("Original files will not be modified. Upload mode sends extracted text and metadata.")
    records: list[dict[str, Any]] = []
    for path in root.rglob("*"):
        if path.is_symlink() or not path.is_file():
            continue
        record = document_record(path, root)
        if record:
            records.append(record)
            print(f"FOUND: {record['path']} ({len(record['text'])} text characters)")

    print(f"\nIndexed locally: {len(records)} document(s).")
    if not records:
        print("No supported files with extractable text were found.")
        return 0

    if not args.upload:
        print("PREVIEW ONLY: nothing was sent to the server.")
        print("Preview complete. Nothing was sent to the server.")
        return 0

    api_url = os.getenv("IQOS_API_URL", "").strip()
    api_key = os.getenv("IQOS_API_KEY", "").strip()
    if not api_url or not api_key:
        print("ERROR: set IQOS_API_URL and IQOS_API_KEY as Windows environment variables first.", file=sys.stderr)
        return 2

    try:
        for start in range(0, len(records), BATCH_SIZE):
            send_batch(api_url, api_key, records[start:start + BATCH_SIZE])
    except (ValueError, RuntimeError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    print("Done. Only metadata and extracted text were sent; original files were not uploaded or modified.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
